"""Builds the availability sheet template Ron edits in Google Sheets.

Run from the repo root:  python _src/make_availability_sheet.py
Upload the .xlsx to Google Sheets, then publish the "Availability" tab as CSV (see README).
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path.home() / "Downloads" / "Reel Irie Availability.xlsx"
STATUSES = ["off", "on call", "booked", "sunset only", "open"]
ROWS = 200

F = "Arial"
HEAD = PatternFill("solid", fgColor="111111")
INPUT = PatternFill("solid", fgColor="FFF9DB")
thin = Side(style="thin", color="D0D0D0")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ---------- Availability (the tab that gets published) ----------
ws = wb.active
ws.title = "Availability"
ws.append(["date", "status", "note"])
for c in ws[1]:
    c.font = Font(name=F, bold=True, color="FFD614", size=12)
    c.fill = HEAD
    c.alignment = Alignment(vertical="center")
ws.row_dimensions[1].height = 24

# Example row: a past date, so it has no effect on the site
ws.append([date(2026, 10, 1), "on call", "Example row. Past dates are ignored, safe to delete."])

for r in range(2, ROWS + 2):
    for col in "ABC":
        cell = ws[f"{col}{r}"]
        cell.font = Font(name=F, size=11, italic=(r == 2), color="777777" if r == 2 else "000000")
        cell.fill = INPUT
        cell.border = BOX
    ws[f"A{r}"].number_format = "mm/dd/yyyy"

ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 56
ws.freeze_panes = "A2"

dv_status = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True,
                           showErrorMessage=True, errorTitle="Pick a status",
                           error="Choose: off, on call, booked, sunset only or open")
dv_status.add(f"B2:B{ROWS + 1}")
dv_date = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True,
                         showErrorMessage=True, errorTitle="Date", error="Enter a date like 10/18/2026")
dv_date.add(f"A2:A{ROWS + 1}")
ws.add_data_validation(dv_status)
ws.add_data_validation(dv_date)

rng = f"A2:C{ROWS + 1}"
for words, color in [(("off", "on call", "booked"), "F8C9C6"), (("sunset only",), "FFE79A"), (("open",), "C8F0D0")]:
    test = "OR(" + ",".join(f'$B2="{w}"' for w in words) + ")"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[test], fill=PatternFill("solid", fgColor=color)))

# ---------- How to use ----------
hw = wb.create_sheet("How to use")
lines = [
    ("Reel Irie Charters: Availability", "title"),
    ("", None),
    ("Add one row per date that's different from the normal week.", None),
    ("Normal week (set on the website): weekdays are Sunset Irie Cruise only, weekends run every trip.", None),
    ("Edit only the yellow cells on the Availability tab. The website updates within a minute.", None),
    ("", None),
    ("status", "head"),
    ("off  /  on call  /  booked", "Blocks the date. Greyed out in the forecast; the form asks for another day."),
    ("sunset only", "Only the Sunset Irie Cruise that day (handy for a weekend when you're short on time)."),
    ("open", "Every trip that day (handy for a weekday you're free)."),
    ("", None),
    ("note", "head"),
    ("Optional", "Shown to customers under the date for sunset only / open days. Kept private for off days."),
    ("", None),
    ("Tips", "head"),
    ("Past dates are ignored, so there's no need to clean up old rows.", None),
    ("Don't rename the Availability tab or its header row (date, status, note).", None),
]
for i, (a, b) in enumerate(lines, start=1):
    hw.cell(row=i, column=1, value=a).font = Font(name=F, size=11)
    if b == "title":
        hw.cell(row=i, column=1).font = Font(name=F, size=16, bold=True)
    elif b == "head":
        hw.cell(row=i, column=1).font = Font(name=F, size=12, bold=True, color="B8860B")
    elif b:
        hw.cell(row=i, column=2, value=b).font = Font(name=F, size=11)
        hw.cell(row=i, column=1).font = Font(name=F, size=11, bold=True)
hw.column_dimensions["A"].width = 30
hw.column_dimensions["B"].width = 90

wb.save(OUT)
print("saved", OUT)
